from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image, ImageDraw, ImageFont
out=Path('deliverables/quality-guide')
im=Image.new('RGB',(1600,630),'white');d=ImageDraw.Draw(im)
font=ImageFont.truetype('C:/Windows/Fonts/calibri.ttf',30)
bold=ImageFont.truetype('C:/Windows/Fonts/calibrib.ttf',31)
d.text((40,12),'MAD (mm)',font=bold,fill='black')
for tick in [0,20,40,60,80,100]:
 x=240+tick*11.5;d.line((x,65,x,555),fill='#e6e8eb',width=2);d.text((x-12,565),str(tick),font=font,fill='black')
for i,(label,v) in enumerate(zip(['0.00','0.10','0.20','0.50','1.00','2.00'],[95,93,91,85,75,55])):
 y=80+i*78;d.text((70,y+2),label,font=font,fill='black');d.rectangle((240,y,240+v*11.5,y+43),fill='#395c76');d.text((253+v*11.5,y+2),f'{v}%',font=font,fill='black')
d.text((630,603),'Initial quality score (%)',font=font,fill='black');im.save(out/'quality-chart.png')
doc=Document();sec=doc.sections[0];sec.page_width=Inches(8.5);sec.page_height=Inches(11);sec.top_margin=sec.bottom_margin=Inches(.7);sec.left_margin=sec.right_margin=Inches(.8)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2']:
 st=doc.styles[name];st.font.name='Calibri';st.font.color.rgb=RGBColor(0,0,0)
normal=doc.styles['Normal'];normal.font.size=Pt(11);normal.paragraph_format.space_after=Pt(7)
doc.styles['Title'].font.size=Pt(25)
doc.styles['Heading 1'].font.size=Pt(16);doc.styles['Heading 1'].paragraph_format.space_before=Pt(12)
doc.styles['Heading 2'].font.size=Pt(12)
def p(text,bold=False):
 par=doc.add_paragraph();par.add_run(text).bold=bold;return par
def eq(text):
 par=doc.add_paragraph();par.paragraph_format.space_after=Pt(9)
 math=OxmlElement('m:oMath');r=OxmlElement('m:r');t=OxmlElement('m:t');t.text=text;r.append(t);math.append(r);par._p.append(math)
def table(headers,rows,widths):
 t=doc.add_table(rows=1,cols=len(headers));t.autofit=False
 for c,w in zip(t.columns,widths):c.width=Inches(w)
 for c,label in zip(t.rows[0].cells,headers):
  c.text=label
  sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'E8EDF1');c._tc.get_or_add_tcPr().append(sh)
  for run in c.paragraphs[0].runs:run.bold=True
 repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
 for row in rows:
  cells=t.add_row().cells
  for c,text in zip(cells,row):c.text=str(text)
 for row in t.rows:
  no=OxmlElement('w:cantSplit');row._tr.get_or_add_trPr().append(no)
  for c in row.cells:
   for par in c.paragraphs:
    par.paragraph_format.space_after=Pt(5);par.paragraph_format.space_before=Pt(5)
 return t

doc.add_paragraph('Ring scan quality percentage explained','Title')
p('How the app calculates the score for each finger and how it differs from the ring size recommendation.')
p('The percentage is an image measurement quality score. It is a heuristic based on measurement consistency, resolution and processing checks. It has not been validated as a percentage of physical accuracy or as a probability that a ring will fit.',True)
doc.add_heading('From photo to quality score',1)
p('Detect finger edges → measure several widths → calculate median width and variation → calculate the initial score → apply quality limits → check a resized photo → display the final percentage.')
doc.add_heading('1  Convert pixels into millimetres',1)
p('The standard reference card is 85.6 mm wide. Its average detected long-edge length gives the image scale. Each accepted measurement line across a finger is converted using that scale.')
eq('Pixels per mm = average card width in pixels ÷ 85.6')
eq('Finger width in mm = line length in pixels ÷ pixels per mm')
p('Example: a card measuring 428 pixels gives 5 pixels per mm. A finger line measuring 100 pixels therefore gives a width of 20 mm.')
doc.add_heading('2  Calculate width consistency',1)
p('The median of the accepted widths is the finger width used by the calculation. MAD means median absolute deviation: calculate the absolute distance of every accepted width from the median, then take the median of those distances.')
eq('MAD = median of |each width − median width|')
table(['Accepted width','Distance from median'],[['19.6 mm','0.4 mm'],['19.8 mm','0.2 mm'],['20.0 mm','0.0 mm'],['20.2 mm','0.2 mm'],['20.4 mm','0.4 mm']],[3.45,3.45])
p('In this example, median width = 20 mm and MAD = 0.2 mm. Smaller MAD means the accepted measurement lines agree more closely.')

doc.add_page_break();doc.add_heading('3  Calculate the initial quality score',1)
p('Let W be the median finger width in mm, D be MAD in mm and S be the image scale in pixels per mm. The current code uses this formula:')
eq('Q = max(20, min(95, round(95 − 400 × D ÷ W),')
eq('                         round(S × 22 + 35)))')
p('The consistency term lowers the score when the widths vary. The resolution term can lower it for a small reference image. The initial score is limited to the range 20% to 95% for a measured finger; an unmeasured finger receives 0%.')
p('Example with W = 20 mm, D = 0.2 mm and S = 5 pixels per mm:')
eq('Consistency term = round(95 − 400 × 0.2 ÷ 20) = 91')
eq('Resolution term = round(5 × 22 + 35) = 145')
eq('Initial quality = max(20, min(95, 91, 145)) = 91%')
doc.add_heading('How variation affects the initial score',1)
p('This chart holds width at 20 mm and assumes enough resolution that the resolution term does not limit the score. These values are before the quality limits on the next page.')
doc.add_picture(str(out/'quality-chart.png'),width=Inches(6.8))
p('MAD values of 0.00, 0.10, 0.20, 0.50, 1.00 and 2.00 mm produce initial scores of 95%, 93%, 91%, 85%, 75% and 55%, respectively.')
p('The constants 95, 400, 22 and 35 are implementation choices. They do not come from a validated study of ring fit or circumference accuracy.',True)

doc.add_page_break();doc.add_heading('4  Apply quality limits',1)
p('A limit can lower the score but cannot increase it. If several conditions apply, the lowest applicable limit wins.')
table(['Condition','Maximum','What it indicates'],[
['Contour-only edges are used','55%','Less certainty in the detected finger boundaries'],
['Measurement is near a joint','70%','The measurement position can affect width'],
['GrabCut fallback segmentation is used','65%','The main segmentation method was unavailable'],
['Second-scale check disagrees or fails','50%','Processing the same photo gives inconsistent results'],
['Finger cannot be measured','0%','No usable finger measurement']],[2.8,.8,3.3])
doc.add_heading('5  Check processing stability',1)
p('When resolution allows, the app measures a comparison image reduced to 85% of the normalized image dimensions. To pass, the second measurement must select the same table size and its circumference must differ by no more than 0.75 mm. A failed comparison or an unusable second measurement limits the score to 50%.')
p('The comparison is skipped when the expected comparison scale is below 3 pixels per mm. A skipped check does not validate stability.')
p('Worked example: an initial 91% becomes 70% if the measurement is near a joint. If the resized-photo check then disagrees, the final displayed score becomes 50%.')
doc.add_heading('How the percentage relates to ring size',1)
p('The ring size is calculated separately from the quality score:')
eq('Estimated circumference = median finger width × π')
eq('Product size = lookup of circumference in your sizing table')
p('The quality percentage does not change the selected product size. The test dropdown values such as +1% and +2% adjust the estimated circumference and recalculate the table size; they are separate from the image quality score.')
p('Manual circumference corrections also leave the photo quality percentage unchanged.')
doc.add_heading('What a high percentage means',1)
p('A high score means the visible width measurements passed the implemented consistency checks. It does not prove that the circumference or selected ring size is correct. The width × π formula assumes a circular finger cross-section, while a single top-view photo does not directly measure finger depth. A repeatable width can therefore still produce the wrong circumference.',True)
p('Implementation basis: hand_worker.py, measurement_stability.py and measurement_trace.py in the local ring-sizing project. Explanation prepared 15 October 2026.')
doc.core_properties.title='Ring scan quality percentage explained';doc.core_properties.subject='Calculation and limitations of the finger scan quality score';doc.core_properties.author='Senbioys'
path=out/'Ring scan quality percentage explained.docx';doc.save(path);print(path.resolve())
